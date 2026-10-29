"""Exporta las métricas a un Excel listo para Power BI (una hoja = una tabla con nombre).

Archivo: salida/metricas_instagram.xlsx. El bot lo regenera una vez por día; en Power BI se carga con
Obtener datos → Excel y se actualiza con el botón "Actualizar".
"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import config
import db
import metricas

ARCHIVO = config.PROJECT_DIR / "salida" / "metricas_instagram.xlsx"
DIAS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]


def _fecha(iso: str | None) -> datetime | None:
    """Fecha local sin zona horaria (Power BI la reconoce como fecha/hora)."""
    if not iso:
        return None
    try:
        return datetime.fromisoformat(iso.replace("+0000", "+00:00")).astimezone(config.ZONA).replace(tzinfo=None)
    except ValueError:
        return None


def _num(valor):
    return None if valor is None else float(valor) if isinstance(valor, float) and not valor.is_integer() else (
        int(valor) if isinstance(valor, (int, float)) else valor)


def _hojas() -> dict[str, tuple[list[str], list[list]]]:
    with db.conexion() as con:
        media = con.execute("SELECT * FROM media WHERE publicada_en IS NOT NULL ORDER BY publicada_en").fetchall()
        diarias = con.execute("SELECT * FROM metricas_diarias ORDER BY fecha").fetchall()
        pubs = con.execute("SELECT * FROM publicaciones ORDER BY id").fetchall()
        banco = con.execute("SELECT * FROM banco_fotos ORDER BY id").fetchall()
        promo = con.execute("SELECT * FROM promo_meses ORDER BY mes").fetchall()
        mensajes = con.execute("SELECT categoria, estado, recibido FROM mensajes ORDER BY recibido").fetchall()

    filas_media = []
    for m in media:
        f = _fecha(m["publicada_en"])
        interacciones = (m["likes"] or 0) + (m["comentarios"] or 0) + (m["guardados"] or 0) + (m["compartidos"] or 0)
        filas_media.append([
            m["id"], f.date() if f else None, f.strftime("%H:%M") if f else None, f.hour if f else None,
            DIAS[f.weekday()] if f else None, (f.weekday() + 1) if f else None, m["tipo"],
            m["likes"], m["comentarios"], _num(m["guardados"]), _num(m["compartidos"]), _num(m["alcance"]),
            _num(m["vistas"]), interacciones,
            round(interacciones / m["alcance"], 4) if m["alcance"] else None,
            (m["caption"] or "").replace("\n", " ")[:300], m["permalink"]])

    horarios = metricas.resumen_horarios()
    with db.conexion() as con:
        todos = con.execute("SELECT publicada_en, likes, comentarios FROM media WHERE publicada_en IS NOT NULL").fetchall()
    acumulado = {}
    for r in todos:
        f = _fecha(r["publicada_en"])
        if f:
            acumulado.setdefault((f.weekday(), f.hour), []).append(r["likes"] + 3 * r["comentarios"])

    return {
        "Publicaciones": (
            ["ID", "Fecha", "Hora", "Hora_numero", "Dia_semana", "Dia_orden", "Tipo", "Likes", "Comentarios",
             "Guardados", "Compartidos", "Alcance", "Vistas", "Interacciones", "Tasa_interaccion", "Descripcion",
             "Link"], filas_media),
        "Seguidores_diarios": (
            ["Fecha", "Seguidores", "Alcance_dia", "Publicaciones_totales"],
            [[datetime.fromisoformat(d["fecha"]).date(), d["seguidores"], d["alcance"], d["publicaciones"]]
             for d in diarias]),
        "Horarios": (
            ["Dia_semana", "Dia_orden", "Hora", "Interaccion_promedio", "Cantidad_publicaciones"],
            [[DIAS[d], d + 1, h, round(sum(v) / len(v), 1), len(v)] for (d, h), v in sorted(acumulado.items())]),
        "Programacion_bot": (
            ["ID", "Tipo", "Estado", "Automatica", "Programada_para", "Tema", "Intentos", "Error"],
            [[p["id"], p["tipo"], p["estado"], "Sí" if p["auto"] else "No", _fecha(p["programada_para"]),
              p["tema"], p["intentos"], p["error"]] for p in pubs]),
        "Banco_fotos": (
            ["ID", "Pie_de_foto", "Destino", "Origen", "Usos_post", "Usos_historia", "Ultimo_uso", "Agregada", "Activa"],
            [[b["id"], b["nombre"], {"historia": "Solo historias", "post": "Solo publicaciones"}.get(b["uso"], "Ambos"),
              b["origen"], b["usos_post"], b["usos_historia"], _fecha(b["ultimo_uso"]), _fecha(b["agregada"]),
              "Sí" if b["activa"] else "No"] for b in banco]),
        "Publicidad": (
            ["Mes", "Presupuesto", "Gastado", "Moneda", "Alcance", "Impresiones", "Conversaciones",
             "Costo_por_conversacion"],
            [[p["mes"], p["presupuesto"], _num(p["gastado"]), p["moneda"], _num(p["alcance"]), _num(p["impresiones"]),
              _num(p["conversaciones"]),
              round(p["gastado"] / p["conversaciones"], 2) if p["gastado"] and p["conversaciones"] else None]
             for p in promo]),
        "Mensajes": (
            ["Fecha", "Categoria", "Estado"],
            [[_fecha(m["recibido"]), m["categoria"], m["estado"]] for m in mensajes]),
        "Resumen": (
            ["Indicador", "Valor"],
            [["Actualizado", datetime.now(config.ZONA).replace(tzinfo=None)],
             ["Publicaciones medidas", len(filas_media)],
             ["Recomendación de horario", horarios["nota"]]]),
    }


def exportar(destino: Path = ARCHIVO) -> Path:
    libro = Workbook()
    libro.remove(libro.active)
    encabezado = PatternFill("solid", fgColor="276347")
    for nombre, (columnas, filas) in _hojas().items():
        hoja = libro.create_sheet(nombre)
        hoja.append(columnas)
        for fila in filas:
            hoja.append(fila)
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = encabezado
        for i, col in enumerate(columnas, start=1):
            letra = get_column_letter(i)
            ancho = max([len(str(col))] + [len(str(f[i - 1])) for f in filas[:200] if f[i - 1] is not None])
            hoja.column_dimensions[letra].width = min(max(ancho + 2, 10), 60)
            for celda in hoja[letra][1:]:
                if isinstance(celda.value, datetime):
                    celda.number_format = "dd/mm/yyyy hh:mm"
                elif hasattr(celda.value, "isoformat") and not isinstance(celda.value, str):
                    celda.number_format = "dd/mm/yyyy"
        hoja.freeze_panes = "A2"
        # Una tabla de Excel por hoja: Power BI la muestra por nombre (tbl_Publicaciones, etc.)
        ultima = f"{get_column_letter(len(columnas))}{max(len(filas), 1) + 1}"
        if not filas:
            hoja.append([None] * len(columnas))
        tabla = Table(displayName=f"tbl_{nombre}", ref=f"A1:{ultima}")
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium7", showRowStripes=True)
        hoja.add_table(tabla)
    destino.parent.mkdir(exist_ok=True)
    temporal = destino.with_suffix(".tmp.xlsx")
    libro.save(temporal)
    try:
        temporal.replace(destino)
    except PermissionError:  # el Excel está abierto: se guarda al lado para no perder la actualización
        alternativo = destino.with_name(destino.stem + "_nuevo.xlsx")
        temporal.replace(alternativo)
        return alternativo
    return destino


if __name__ == "__main__":
    print(f"Excel generado: {exportar()}")
